import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

STATUS_LABELS = {
    "SCHEDULED": "Agendado",
    "ON_YARD": "No pátio",
    "IN_OPERATION": "Em operação",
    "AWAITING_CONCLUSION": "Aguardando NF",
    "FINISHED": "Finalizado",
    "CANCELLED": "Cancelado",
}

TEXT_HEADERS = [
    "Placa",
    "Empresa",
    "Motorista",
    "Produto",
    "Tipo de carga",
    "Área",
    "Status",
    "Operação",
]

TIMESTAMPS = [
    ("Criado em", "created_at"),
    ("Chegada", "arrival_time"),
    ("Início operação", "start_time"),
    ("Aguardando NF", "awaiting_conclusion_time"),
    ("Fim", "end_time"),
]

DURATIONS = [
    ("Tempo até chegada (min)", "created_at", "arrival_time"),
    ("Tempo em espera (min)", "arrival_time", "start_time"),
    ("Tempo em operação (min)", "start_time", "awaiting_conclusion_time"),
    ("Tempo aguardando NF (min)", "awaiting_conclusion_time", "end_time"),
    ("Tempo total (min)", "created_at", "end_time"),
]

DATE_FORMAT = "DD/MM/YYYY"
TIME_FORMAT = "HH:MM"  # 24h
MINUTES_FORMAT = "0.0"


def _to_local_naive(dt):
    if dt is None:
        return None
    return timezone.localtime(dt).replace(tzinfo=None, microsecond=0)


def _minutes_between(start, end):
    if start is None or end is None:
        return None
    seconds = (end - start).total_seconds()
    if seconds < 0:
        return None
    return round(seconds / 60, 1)


def _column_layout():
    headers = list(TEXT_HEADERS)
    formats = [None] * len(TEXT_HEADERS)
    widths = [12, 24, 26, 18, 14, 22, 16, 12]

    for label, _ in TIMESTAMPS:
        headers += [f"{label} (data)", f"{label} (hora)"]
        formats += [DATE_FORMAT, TIME_FORMAT]
        widths += [14, 10]

    for label, _, _ in DURATIONS:
        headers.append(label)
        formats.append(MINUTES_FORMAT)
        widths.append(18)

    return headers, formats, widths


def build_queue_entries_workbook(entries) -> bytes:
    headers, formats, widths = _column_layout()

    wb = Workbook()
    ws = wb.active
    ws.title = "Histórico"
    ws.append(headers)

    for entry in entries.iterator(chunk_size=500):
        row = [
            entry.truck_plate,
            entry.company_name or "",
            entry.truck_driver,
            entry.truck_product,
            entry.truck_cargo_type,
            entry.area.name if entry.area else "",
            STATUS_LABELS.get(entry.status, entry.status),
            entry.job or "",
        ]

        for _, field in TIMESTAMPS:
            local = _to_local_naive(getattr(entry, field))
            row += [local.date() if local else None, local.time() if local else None]

        for _, start_field, end_field in DURATIONS:
            row.append(
                _minutes_between(getattr(entry, start_field), getattr(entry, end_field))
            )

        ws.append(row)
        row_idx = ws.max_row
        for col_idx, fmt in enumerate(formats, start=1):
            if fmt:
                ws.cell(row=row_idx, column=col_idx).number_format = fmt

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
